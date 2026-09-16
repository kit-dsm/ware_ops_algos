"""
Enumeriert alle Pick-Kombinationen ueber M Gassen und berechnet pro Kombination:
  - Intervallstruktur via interval_builder
  - optimales wait* (Python-Minimierung auf J(wait))
  - kritischen Schwellwert delta* (J'(0)=0), falls vorhanden

Ergebnisse werden als XLSX exportiert (Summary + Details).
"""

from __future__ import annotations

import argparse
import math
import os
import sys
from dataclasses import dataclass
from datetime import datetime
from math import factorial
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from scipy.optimize import minimize_scalar

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

try:
    from . import interval_builder as ib
    from .models import WarehouseInstance
    from .validation import compute_J_for_wait, find_critical_delta
except ImportError:
    import interval_builder as ib
    from models import WarehouseInstance
    from validation import compute_J_for_wait, find_critical_delta


@dataclass(frozen=True)
class LayoutParams:
    M: int = 8
    N_L: int = 17
    w: float = 1.0
    L: float = 17.0
    v: float = 1.0
    t_p: float = 0.0


@dataclass(frozen=True)
class SolveParams:
    delta_exp_for_wait_opt: float = 28.8
    n_picks: int = 3
    n_picks_after: int = 3
    wait_upper_bound: float = 200.0
    delta_lo: float = 1.0
    delta_hi: float = 500.0
    delta_tol: float = 1e-6


def get_aisle_combinations(n_picks: int, n_aisles: int) -> List[Tuple[Tuple[int, ...], float]]:
    """Multinomial-Verteilungen von n_picks Picks auf n_aisles Gassen."""
    results: List[Tuple[Tuple[int, ...], float]] = []

    def _recurse(remaining: int, idx: int, current: List[int]) -> None:
        if idx == n_aisles:
            if remaining == 0:
                combo = tuple(current)
                coeff = factorial(n_picks)
                for k in combo:
                    coeff //= factorial(k)
                prob = coeff * (1.0 / n_aisles) ** n_picks
                results.append((combo, prob))
            return
        for k in range(remaining + 1):
            _recurse(remaining - k, idx + 1, current + [k])

    _recurse(n_picks, 0, [])
    return results


def combo_to_instance(combo: Tuple[int, ...], layout: LayoutParams) -> WarehouseInstance:
    visited = [i + 1 for i, n in enumerate(combo) if n > 0]
    n_list = [combo[a - 1] for a in visited]
    return WarehouseInstance(
        M=layout.M,
        N_L=layout.N_L,
        w=layout.w,
        L=layout.L,
        A=visited,
        n_list=n_list,
        v=layout.v,
        t_p=layout.t_p,
        P=[],
        arrival_times=[],
    )


def _safe_float(value: Optional[float]) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return float(value)


def _to_float(value: object, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def solve_combination(
    combo: Tuple[int, ...],
    probability: float,
    layout: LayoutParams,
    solve: SolveParams,
) -> Dict[str, object]:
    inst = combo_to_instance(combo, layout)
    intervals, t_end = ib.build_interval_data(inst)

    lam = 1.0 / solve.delta_exp_for_wait_opt

    def j_py(wait: float) -> float:
        if wait < 0:
            return float("inf")
        return compute_J_for_wait(
            lam,
            wait,
            intervals,
            t_end,
            n_picks_after=solve.n_picks_after,
        )

    opt = minimize_scalar(j_py, bounds=(0.0, solve.wait_upper_bound), method="bounded")
    wait_opt = float(opt.x)
    if wait_opt < 1e-5:
        wait_opt = 0.0
    j_opt = float(opt.fun)

    delta_star: Optional[float]
    delta_error = ""
    try:
        delta_star = find_critical_delta(
            intervals,
            t_end,
            n_picks_after=solve.n_picks_after,
            delta_lo=solve.delta_lo,
            delta_hi=solve.delta_hi,
            tol=solve.delta_tol,
        )
    except Exception as exc:  # bewusst robust fuer alle Kombinationen
        delta_star = None
        delta_error = str(exc)

    ed_i0 = float(intervals[0].ed_coeffs[0]) if intervals else 0.0

    return {
        "combo": str(combo),
        "probability": probability,
        "visited": str(inst.A),
        "n_list": str(inst.n_list),
        "k": inst.k,
        "T_end": float(t_end),
        "n_intervals": len(intervals),
        "ed_i0": ed_i0,
        "wait_opt": wait_opt,
        "J_opt": j_opt,
        "delta_star": _safe_float(delta_star),
        "lambda_star": _safe_float((1.0 / delta_star) if delta_star else None),
        "delta_status": "ok" if delta_star is not None else "no-root",
        "delta_error": delta_error,
    }


def _excel_value(value):
    if isinstance(value, float):
        if math.isfinite(value):
            return value
        if math.isnan(value):
            return "nan"
        return "inf" if value > 0 else "-inf"
    return value


def _autosize_worksheet_columns(ws) -> None:
    from openpyxl.utils import get_column_letter

    for col_idx, column_cells in enumerate(
        ws.iter_cols(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column),
        start=1,
    ):
        max_len = 0
        for cell in column_cells:
            text = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(text))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(8, max_len + 2)


def export_results_xlsx(details: Sequence[Dict[str, object]], summary: Dict[str, object], output_path: Path) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("XLSX-Export nicht moeglich: openpyxl fehlt. pip install openpyxl") from exc

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "summary"

    ws_summary.append(["metric", "value"])
    for key, value in summary.items():
        ws_summary.append([key, _excel_value(value)])
    _autosize_worksheet_columns(ws_summary)

    ws_details = wb.create_sheet(title="details")
    if details:
        headers = list(details[0].keys())
        ws_details.append(headers)
        for row in details:
            ws_details.append([_excel_value(row.get(h)) for h in headers])

        for col_idx, header in enumerate(headers, start=1):
            for r_idx in range(2, ws_details.max_row + 1):
                cell = ws_details.cell(row=r_idx, column=col_idx)
                if isinstance(cell.value, float) and math.isfinite(cell.value):
                    if header in {"probability"}:
                        cell.number_format = "0.000000"
                    else:
                        cell.number_format = "0.000000"
        _autosize_worksheet_columns(ws_details)
    else:
        ws_details.append(["Hinweis", "Keine Detaildaten vorhanden"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))


def _default_output_path() -> Path:
    root = Path(__file__).resolve().parents[3]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return root / "Results" / f"all_combinations_wait_delta_{ts}.xlsx"


def run_all_combinations(
    layout: LayoutParams,
    solve: SolveParams,
    max_combos: Optional[int] = None,
    verbose: bool = True,
) -> Tuple[List[Dict[str, object]], Dict[str, object]]:
    combos = get_aisle_combinations(solve.n_picks, layout.M)
    combos = sorted(combos, key=lambda x: x[1], reverse=True)
    if max_combos is not None:
        combos = combos[:max_combos]

    if verbose:
        print(f"Kombinationen zu berechnen: {len(combos)}")

    details: List[Dict[str, object]] = []
    for idx, (combo, prob) in enumerate(combos, start=1):
        row = solve_combination(combo, prob, layout, solve)
        details.append(row)
        if verbose and (idx % 20 == 0 or idx == len(combos)):
            print(
                f"[{idx:>4}/{len(combos)}] combo={row['combo']} "
                f"wait*={row['wait_opt']:.4f} J*={row['J_opt']:.4f} "
                f"delta*={row['delta_star']}"
            )

    p_sum = sum(_to_float(r.get("probability")) for r in details)
    e_wait = sum(_to_float(r.get("probability")) * _to_float(r.get("wait_opt")) for r in details)
    e_j = sum(_to_float(r.get("probability")) * _to_float(r.get("J_opt")) for r in details)

    delta_rows = [r for r in details if r["delta_star"] is not None]
    p_delta = sum(_to_float(r.get("probability")) for r in delta_rows)
    e_delta_cond = (
        sum(_to_float(r.get("probability")) * _to_float(r.get("delta_star")) for r in delta_rows) / p_delta
        if p_delta > 0
        else None
    )
    e_delta_weighted = sum(_to_float(r.get("probability")) * _to_float(r.get("delta_star")) for r in delta_rows)

    summary: Dict[str, object] = {
        "n_combos": len(details),
        "n_delta_star_ok": len(delta_rows),
        "probability_sum": p_sum,
        "probability_sum_delta_ok": p_delta,
        "E_wait_opt": e_wait,
        "E_J_opt": e_j,
        "E_delta_star_weighted": e_delta_weighted,
        "E_delta_star_conditional": e_delta_cond if e_delta_cond is not None else "n/a",
        "delta_exp_for_wait_opt": solve.delta_exp_for_wait_opt,
        "n_picks": solve.n_picks,
        "n_picks_after": solve.n_picks_after,
        "M": layout.M,
        "N_L": layout.N_L,
        "L": layout.L,
    }
    return details, summary


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Berechnet fuer alle Gassenkombinationen (gewichtete Multinomial-Wahrscheinlichkeit) "
            "das optimale wait* und den kritischen Schwellwert delta* und exportiert XLSX."
        )
    )
    parser.add_argument("--delta-exp", type=float, default=28.8,
                        help="delta fuer wait*-Optimierung (lambda=1/delta)")
    parser.add_argument("--n-picks", type=int, default=3,
                        help="Anzahl Picks pro Kombination (Multinomial-Enumeration)")
    parser.add_argument("--n-picks-after", type=int, default=3,
                        help="n_picks_after fuer ECostAfter in der Kostenfunktion")
    parser.add_argument("--M", type=int, default=8, help="Anzahl Gassen")
    parser.add_argument("--N-L", dest="N_L", type=int, default=17, help="Aisle node count")
    parser.add_argument("--L", type=float, default=17.0, help="Aisle length")
    parser.add_argument("--w", type=float, default=1.0, help="Aisle spacing")
    parser.add_argument("--v", type=float, default=1.0, help="Speed")
    parser.add_argument("--t-p", type=float, default=0.0, help="Pick time")
    parser.add_argument("--wait-upper", type=float, default=200.0,
                        help="Obere Grenze fuer die numerische wait*-Suche")
    parser.add_argument("--delta-lo", type=float, default=1.0,
                        help="Untere Schranke fuer delta* Root-Suche")
    parser.add_argument("--delta-hi", type=float, default=500.0,
                        help="Obere Schranke fuer delta* Root-Suche")
    parser.add_argument("--delta-tol", type=float, default=1e-6,
                        help="Toleranz fuer delta* Root-Suche")
    parser.add_argument("--max-combos", type=int, default=None,
                        help="Optional: nur die ersten N (nach Wahrscheinlichkeit sortiert)")
    parser.add_argument("--output", type=str, default=None,
                        help="Pfad zur XLSX-Datei")
    parser.add_argument("--quiet", action="store_true", help="Weniger Konsolenausgabe")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    layout = LayoutParams(M=args.M, N_L=args.N_L, w=args.w, L=args.L, v=args.v, t_p=args.t_p)
    solve = SolveParams(
        delta_exp_for_wait_opt=args.delta_exp,
        n_picks=args.n_picks,
        n_picks_after=args.n_picks_after,
        wait_upper_bound=args.wait_upper,
        delta_lo=args.delta_lo,
        delta_hi=args.delta_hi,
        delta_tol=args.delta_tol,
    )

    details, summary = run_all_combinations(
        layout=layout,
        solve=solve,
        max_combos=args.max_combos,
        verbose=not args.quiet,
    )

    output = Path(args.output) if args.output else _default_output_path()
    export_results_xlsx(details, summary, output)

    print("\nFertig.")
    print(f"XLSX: {output}")
    print(f"E[wait*] = {summary['E_wait_opt']:.6f}")
    print(f"E[J*]    = {summary['E_J_opt']:.6f}")
    print(f"E[delta*] (gewichtet) = {summary['E_delta_star_weighted']:.6f}")


if __name__ == "__main__":
    main()


