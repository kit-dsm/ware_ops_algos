from pathlib import Path
import math
from typing import Dict, List, Sequence, Tuple, Union

try:
    from .io_cli import parse_args
    from .models import BatchComparisonResult
    from .order_optimizer import analyze_batches, batch_sample_rows, batch_summary_rows
except ImportError:  # Direktes Ausführen als Skript
    from io_cli import parse_args
    from models import BatchComparisonResult
    from order_optimizer import analyze_batches, batch_sample_rows, batch_summary_rows


Row = Dict[str, Union[float, int, str]]
DEFAULT_BATCH_SUMMARY_DISCRETE_SHEET = "Batch Summary Discrete"
DEFAULT_BATCH_SUMMARY_CONTINUOUS_SHEET = "Batch Summary Continuous"
DEFAULT_BATCH_SAMPLES_DISCRETE_SHEET = "Batch Samples Discrete"
DEFAULT_BATCH_SAMPLES_CONTINUOUS_SHEET = "Batch Samples Continuous"
DEFAULT_COMPARE_SHEET = "Engine Compare"
DEFAULT_KPI_COMPARE_SHEET = "KPI Vergleich"
DEFAULT_E_D_DEVIATION_SHEET = "E(D) Abweichungsanalyse"


COMPACT_SAMPLE_COLUMNS: List[str] = [
    "batch", "base_orders", "k",
    "t", "x", "y", "phase", "return_flag",
    "E[D]", "d_rb", "d_pass", "d_front",
    "p1", "p2", "p3", "p4", "p_sum",
    "|V|", "|R|", "|Gp|", "|Gf|",
]


KPI_COMPARE_FIELDS: List[Tuple[str, str]] = [
    ("E(D)",  "E[D]"),
    ("d_rb",  "d_rb"),
    ("d_pass","d_pass"),
    ("d_front","d_front"),
    ("p1",    "p1"),
    ("p2",    "p2"),
    ("p3",    "p3"),
    ("p4",    "p4"),
    ("p_sum", "p_sum"),
    ("Gp",    "|Gp|"),
    ("Gf",    "|Gf|"),
]


def _default_batch_xlsx_path(results_dir: Union[str, Path]) -> Path:
    return Path(results_dir) / "batch_analysis.xlsx"


# def _tag_rows(rows: Sequence[Row], engine_name: str) -> List[Row]:
#     tagged: List[Row] = []
#     for row in rows:
#         tagged_row = dict(row)
#         tagged_row["engine"] = engine_name
#         tagged.append(tagged_row)
#     return tagged


def _extract_k(result: BatchComparisonResult) -> int:
    aisles = {aisle for pos in result.base_order_positions for aisle, _y in pos}
    return len(aisles)


def _build_batch_to_k(results):
    return {r.batch_index: _extract_k(r) for r in results}


def _enrich_rows_with_k(rows, batch_to_k):
    enriched = []
    for row in rows:
        k_value = batch_to_k.get(row.get("batch"), "")
        anchor = "base_orders" if "base_orders" in row else ("batch" if "batch" in row else None)
        if anchor is None:
            new_row: Row = {"k": k_value, **row}
        else:
            new_row = {}
            for key, value in row.items():
                new_row[key] = value
                if key == anchor:
                    new_row["k"] = k_value
        enriched.append(new_row)
    return enriched


def _filter_columns(rows, columns):
    return [{col: row.get(col, "") for col in columns} for row in rows]


def _column_mean(rows, key):
    values = []
    for row in rows:
        v = row.get(key)
        if isinstance(v, bool):
            continue
        if isinstance(v, (int, float)) and not (isinstance(v, float) and math.isnan(v)):
            values.append(float(v))
    return sum(values) / len(values) if values else 0.0


def _build_kpi_comparison_rows(discrete_rows, continuous_rows) -> List[Row]:
    rows: List[Row] = []
    for label, key in KPI_COMPARE_FIELDS:
        d = _column_mean(discrete_rows, key)
        c = _column_mean(continuous_rows, key)
        abs_diff = d - c
        rel_diff = abs_diff / d if d != 0 else 0.0
        rows.append({
            "Vergleichswert": label,
            "Diskret": d,
            "Kontinuierlich": c,
            "absolute Abweichung": abs_diff,
            "relative Abweichung": rel_diff,
        })
    return rows


def _build_e_d_deviation_rows(
    discrete_rows, continuous_rows,
    batch_to_k_discrete, batch_to_k_continuous,
) -> List[Row]:
    """Pro (batch, t) eine Zeile: E[D] diskret vs kontinuierlich + abs Diff, sortiert nach |Diff|."""
    cont_by_key: Dict[Tuple, Row] = {(r.get("batch"), r.get("t")): r for r in continuous_rows}
    out: List[Row] = []
    for d_row in discrete_rows:
        b = d_row.get("batch")
        t = d_row.get("t")
        c_row = cont_by_key.get((b, t))
        if c_row is None:
            continue
        try:
            ed_d = float(d_row.get("E[D]", 0.0))
            ed_c = float(c_row.get("E[D]", 0.0))
            abs_diff = abs(ed_d - ed_c)
        except (TypeError, ValueError):
            ed_d, ed_c, abs_diff = d_row.get("E[D]", ""), c_row.get("E[D]", ""), 0.0
        out.append({
            "batch": b,
            "base_orders_diskret": d_row.get("base_orders", ""),
            "base_orders_kont":    c_row.get("base_orders", ""),
            "k_diskret":           batch_to_k_discrete.get(b, ""),
            "k_kont":              batch_to_k_continuous.get(b, ""),
            "t": t,
            "x": d_row.get("x", ""),
            "y": d_row.get("y", ""),
            "phase": d_row.get("phase", ""),
            "return_flag": d_row.get("return_flag", ""),
            "E[D] diskret": ed_d,
            "E[D] kont":    ed_c,
            "abs Diff":     abs_diff,
        })
    out.sort(key=lambda r: r["abs Diff"] if isinstance(r["abs Diff"], (int, float)) else 0.0, reverse=True)
    return out


def _build_engine_compare_rows(
    discrete_results: Sequence[BatchComparisonResult],
    continuous_results: Sequence[BatchComparisonResult],
) -> List[Row]:
    cont_by_batch = {result.batch_index: result for result in continuous_results}
    rows: List[Row] = []
    for d_res in discrete_results:
        c_res = cont_by_batch.get(d_res.batch_index)
        if c_res is None:
            continue
        rows.append(
            {
                "batch": d_res.batch_index,
                "k": _extract_k(d_res),
                "discrete_actual_case": d_res.actual_case,
                "continuous_actual_case": c_res.actual_case,
                "discrete_actual_detour": d_res.actual_detour,
                "continuous_actual_detour": c_res.actual_detour,
                "delta_actual_detour": c_res.actual_detour - d_res.actual_detour,
                "discrete_mean_actual_oct": d_res.mean_actual_order_completion_time,
                "continuous_mean_actual_oct": c_res.mean_actual_order_completion_time,
                "delta_mean_actual_oct": c_res.mean_actual_order_completion_time - d_res.mean_actual_order_completion_time,
                "discrete_phase1_pred_completion": d_res.phase1_predicted_completion_time,
                "continuous_phase1_pred_completion": c_res.phase1_predicted_completion_time,
                "delta_phase1_pred_completion": c_res.phase1_predicted_completion_time - d_res.phase1_predicted_completion_time,
            }
        )
    return rows


def _excel_value(value: Union[float, int, str]) -> Union[float, int, str]:
    """Normalisiert Zellwerte fuer XLSX ohne zusaetzliches Runden von Floats."""
    if isinstance(value, float):
        if math.isfinite(value):
            return value
        if math.isnan(value):
            return "nan"
        return "inf" if value > 0 else "-inf"
    return value


def _autosize_worksheet_columns(worksheet) -> None:
    """Setzt Spaltenbreiten passend zum sichtbaren Zellinhalt."""
    from openpyxl.utils import get_column_letter

    for col_idx, column_cells in enumerate(
        worksheet.iter_cols(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column),
        start=1,
    ):
        max_len = 0
        for cell in column_cells:
            value = cell.value
            if value is None:
                text = ""
            elif isinstance(value, float) and math.isfinite(value):
                text = str(value)
            else:
                text = str(value)
            if len(text) > max_len:
                max_len = len(text)
        worksheet.column_dimensions[get_column_letter(col_idx)].width = max(8, max_len + 2)


def export_multi_sheet_to_xlsx(sheet_rows, output_path):
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("XLSX-Export nicht moeglich: openpyxl fehlt. pip install openpyxl") from exc

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, rows in sheet_rows:
        ws = wb.create_sheet(title=(sheet_name or "sheet")[:31])
        if not rows:
            ws.append(["Hinweis", "Keine Daten fuer den Export vorhanden"])
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([_excel_value(row[h]) for h in headers])

        for col_idx, header in enumerate(headers, start=1):
            for r_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                if isinstance(cell.value, float) and math.isfinite(cell.value):
                    cell.number_format = "0.00%" if header == "relative Abweichung" else "0.0000"

        _autosize_worksheet_columns(ws)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\nXLSX gespeichert: {out}")


def run_batch_analysis(args, compact: bool = False) -> List[BatchComparisonResult]:
    """Fuehrt die Batchanalyse fuer beide Engines aus und exportiert den Vergleich."""
    discrete_results = analyze_batches(
        args.orders_json,
        args.mapping_json,
        M=args.M,
        N_L=args.N_L,
        w=args.w,
        L=args.L,
        v=args.v,
        t_p=args.t_p,
        known_size=args.batch_known_size,
        total_size=args.batch_total_size,
        max_batches=args.max_batches,
        engine="discrete",
    )
    continuous_results = analyze_batches(
        args.orders_json,
        args.mapping_json,
        M=args.M,
        N_L=args.N_L,
        w=args.w,
        L=args.L,
        v=args.v,
        t_p=args.t_p,
        known_size=args.batch_known_size,
        total_size=args.batch_total_size,
        max_batches=args.max_batches,
        engine="continuous",
    )

    # k pro Batch fuer beide Engines bestimmen
    batch_to_k_d = _build_batch_to_k(discrete_results)
    batch_to_k_c = _build_batch_to_k(continuous_results)

    # k in jede Zeile einfuegen (nach base_orders)
    summary_d = _enrich_rows_with_k(batch_summary_rows(discrete_results), batch_to_k_d)
    summary_c = _enrich_rows_with_k(batch_summary_rows(continuous_results), batch_to_k_c)
    samples_d = _enrich_rows_with_k(batch_sample_rows(discrete_results), batch_to_k_d)
    samples_c = _enrich_rows_with_k(batch_sample_rows(continuous_results), batch_to_k_c)

    output_path = _default_batch_xlsx_path(args.results_dir)

    if compact:
        kpi_rows = _build_kpi_comparison_rows(samples_d, samples_c)
        deviation_rows = _build_e_d_deviation_rows(samples_d, samples_c, batch_to_k_d, batch_to_k_c)
        samples_d = _filter_columns(samples_d, COMPACT_SAMPLE_COLUMNS)
        samples_c = _filter_columns(samples_c, COMPACT_SAMPLE_COLUMNS)
        sheet_rows = [
            (DEFAULT_BATCH_SAMPLES_DISCRETE_SHEET, samples_d),
            (DEFAULT_BATCH_SAMPLES_CONTINUOUS_SHEET, samples_c),
            (DEFAULT_KPI_COMPARE_SHEET, kpi_rows),
            (DEFAULT_E_D_DEVIATION_SHEET, deviation_rows),
        ]
    else:
        compare_rows = _build_engine_compare_rows(discrete_results, continuous_results)
        sheet_rows = [
            (DEFAULT_BATCH_SUMMARY_DISCRETE_SHEET, summary_d),
            (DEFAULT_BATCH_SUMMARY_CONTINUOUS_SHEET, summary_c),
            (DEFAULT_BATCH_SAMPLES_DISCRETE_SHEET, samples_d),
            (DEFAULT_BATCH_SAMPLES_CONTINUOUS_SHEET, samples_c),
            (DEFAULT_COMPARE_SHEET, compare_rows),
        ]

    export_multi_sheet_to_xlsx(sheet_rows, str(output_path))

    print(
        f"Ausgewertete Batches: discrete={len(discrete_results)}, "
        f"continuous={len(continuous_results)} (compact={compact})"
    )
    return discrete_results


def main() -> None:
    args = parse_args()
    try:
        run_batch_analysis(args, compact=True)
    except RuntimeError as exc:
        print(f"\nWarnung: {exc}")


if __name__ == "__main__":
    main()

